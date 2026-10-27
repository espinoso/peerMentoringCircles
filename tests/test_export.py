from io import BytesIO

from openpyxl import load_workbook

from matching.export import HEADERS, build_xlsx
from matching.parser import Participant
from matching.prompt import Group


def test_xlsx_has_circle_sections():
    people = [Participant(id=i, name=f"P{i}", email=f"p{i}@x.org") for i in range(1, 6)]
    groups = [
        Group(name="Circle A", member_ids=[1, 2, 3], rationale="shared topic"),
        Group(name="Circle B", member_ids=[4, 5]),
    ]

    ws = load_workbook(BytesIO(build_xlsx(groups, people))).active
    rows = [[c.value for c in r] for r in ws.iter_rows()]

    assert rows[0] == HEADERS
    # Circle A: name (+ rationale) on first member row, members beneath, no name repeated
    assert rows[1][0].startswith("Circle A") and "shared topic" in rows[1][0]
    assert [r[1] for r in rows[1:4]] == ["P1", "P2", "P3"]
    assert [r[0] for r in rows[2:4]] == [None, None]
    # blank separator row
    assert all(v is None for v in rows[4])
    # Circle B follows the blank row
    assert rows[5][0] == "Circle B"
    assert [r[1] for r in rows[5:7]] == ["P4", "P5"]
    assert rows[6][0] is None
